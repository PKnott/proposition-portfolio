#!/usr/bin/env python3
"""Write gathered bookmaker prices into a daily odds_input workbook.

The other half of `/fill-odds`. The skill does the browsing, the matching and the
judgement; this script does none of it. It takes a JSON file of already-matched
prices and writes them into the book columns of each fixture sheet, paired to
column A by *exact* proposition text.

Splitting it this way is the point. Gathering odds is a job for a browser and a
model, and neither is reproducible; writing 832 numbers into the right cells is a
job for a script, and it has to be. Re-running this is safe: every book column is
re-derived from whatever JSON it is handed, so a second pass with a fuller file
simply overwrites the first.

    python scripts/fill_odds.py --file Outputs/odds_input_2026-08-19.xlsx \
                                --odds odds_gathered.json

odds.json schema -- sheet code, then the proposition text exactly as it appears in
column A, then one key per book from `staking.BOOKS`:

{
  "SP1-01": {
    "Goals - Atletico Madrid - Over 0.5": {"b365": 1.05, "paddypower": 1.04, "skybet": 1.06},
    "Shots - Atletico Madrid - Over 10.5": {"skybet": 1.88}
  }
}

Omissions are meaningful and there are two kinds. A *sheet* left out is skipped
entirely, its cells untouched -- which is what makes it safe to fill the workbook
in batches across a day as more markets are priced up. A *proposition* left out
of a sheet that is present gets 0, the workbook's spelling of "no book is pricing
this line". You never have to write a 0 yourself.

The sheets to fill are read from the workbook's own "Contents" tab, so the
fixture list can change daily without touching this file, and the books are read
from `staking.BOOKS`, so the columns cannot drift out of step with the form.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

import openpyxl

from fpp.report.odds_form import FIRST_BOOK_COL, FIRST_DATA_ROW
from fpp.staking import BOOK_COLUMNS, BOOKS, NOT_OFFERED

# Decimal odds are a multiplier on a stake and so are always above 1.00. The form
# refuses anything in (0, 1.00] at entry; this refuses it on the way in from JSON,
# because a price of 0.85 is not a small error -- it reads downstream as a
# negative edge, gets dropped at the `e >= 1` step, and looks exactly like a line
# that was priced correctly and rejected on its merits.
MIN_DECIMAL_PRICE = 1.0


def get_target_sheets(wb) -> list[str]:
    """The fixture sheets listed on the Contents tab, in order."""
    if "Contents" not in wb.sheetnames:
        raise ValueError("No 'Contents' sheet found in workbook.")
    ws = wb["Contents"]

    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        if row[0].value == "Sheet":
            header_row = row[0].row
            break
    if header_row is None:
        raise ValueError("Could not find the 'Sheet' header row on the Contents tab.")

    sheets = []
    r = header_row + 1
    while True:
        val = ws.cell(row=r, column=1).value
        if val is None or str(val).strip() == "":
            break
        sheets.append(str(val).strip())
        r += 1
    return sheets


def _price(entry: dict, book: str, where: str, bad: list[str]) -> float:
    """One book's price for one proposition, or 0 if it is not pricing it."""
    v = entry.get(book)
    if v is None:
        return NOT_OFFERED
    try:
        v = float(v)
    except (TypeError, ValueError):
        bad.append(f"{where} [{book}] = {v!r} (not a number)")
        return NOT_OFFERED
    if v == NOT_OFFERED:
        return NOT_OFFERED
    if v <= MIN_DECIMAL_PRICE:
        bad.append(f"{where} [{book}] = {v} (decimal odds are above 1.00)")
        return NOT_OFFERED
    return v


def fill_sheet(ws, sheet_odds: dict, sheet_name: str,
               missing_log: list[str], bad: list[str]) -> tuple[int, int]:
    """Fill every book column on one fixture sheet. Returns (filled, unmatched)."""
    filled = missing = 0
    row = FIRST_DATA_ROW
    while True:
        prop = ws.cell(row=row, column=1).value
        if prop is None or str(prop).strip() == "":
            break
        prop = str(prop).strip()
        entry = sheet_odds.get(prop)

        if entry is None:
            prices = [NOT_OFFERED] * len(BOOK_COLUMNS)
            missing += 1
            missing_log.append(f"{sheet_name} | row {row} | {prop}")
        else:
            where = f"{sheet_name} row {row} ({prop})"
            prices = [_price(entry, b, where, bad) for b in BOOK_COLUMNS]
            # A proposition present in the JSON but priced by nobody is no better
            # than an absent one, and counting it as filled would flatter the
            # summary at exactly the moment it should be raising a question.
            if any(p != NOT_OFFERED for p in prices):
                filled += 1
            else:
                missing += 1
                missing_log.append(f"{sheet_name} | row {row} | {prop}  (supplied, no price)")

        for j, price in enumerate(prices):
            ws.cell(row=row, column=FIRST_BOOK_COL + j).value = price
        row += 1
    return filled, missing


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--file", required=True, help="Path to today's odds_input xlsx")
    ap.add_argument("--odds", required=True, help="Path to JSON file of gathered odds")
    ap.add_argument("--no-backup", action="store_true",
                    help="Skip writing a .bak copy before overwriting")
    args = ap.parse_args()

    xlsx_path, odds_path = Path(args.file), Path(args.odds)
    if not xlsx_path.exists():
        sys.exit(f"File not found: {xlsx_path}")
    if not odds_path.exists():
        sys.exit(f"Odds JSON not found: {odds_path}")

    odds_data = json.loads(odds_path.read_text())
    unknown_books = {b for s in odds_data.values() for e in s.values() for b in e} - set(BOOK_COLUMNS)
    if unknown_books:
        sys.exit(f"Unknown book key(s) {sorted(unknown_books)}; "
                 f"this workbook's books are {list(BOOK_COLUMNS)}.")

    # Working files go to `Outputs/Data/`, not next to the form. The form's own
    # folder is meant to hold three files you open, and a `.bak` beside it also
    # made `latest_form` glob over a second `odds_input_*` name.
    work = _work_dir(xlsx_path)

    if not args.no_backup:
        backup_path = work / (xlsx_path.name + ".bak")
        shutil.copy2(xlsx_path, backup_path)
        print(f"Backup written: {backup_path}")

    wb = openpyxl.load_workbook(xlsx_path)
    target_sheets = get_target_sheets(wb)

    missing_log: list[str] = []
    bad: list[str] = []
    total_filled = total_missing = 0
    skipped: list[str] = []

    for sheet_name in target_sheets:
        if sheet_name not in wb.sheetnames:
            skipped.append(f"{sheet_name} (not in workbook)")
            continue
        sheet_odds = odds_data.get(sheet_name)
        if sheet_odds is None:
            skipped.append(f"{sheet_name} (no odds supplied)")
            continue
        filled, missing = fill_sheet(wb[sheet_name], sheet_odds, sheet_name, missing_log, bad)
        total_filled += filled
        total_missing += missing
        print(f"{sheet_name}: {filled} filled, {missing} unmatched -> 0")

    if bad:
        sys.exit("\nRefusing to write; the odds JSON contains impossible prices:\n  "
                 + "\n  ".join(bad[:20])
                 + (f"\n  (+{len(bad) - 20} more)" if len(bad) > 20 else ""))

    wb.save(xlsx_path)

    priced_of = total_filled + total_missing
    print("\n--- Summary ---")
    print(f"Books written: {', '.join(BOOKS.values())}")
    print(f"Sheets processed: {len(target_sheets) - len(skipped)} / {len(target_sheets)}")
    print(f"Propositions filled: {total_filled}"
          + (f" of {priced_of} ({total_filled / priced_of:.0%})" if priced_of else ""))
    print(f"Propositions defaulted to 0 (no bookmaker match): {total_missing}")
    if skipped:
        print(f"Sheets skipped entirely: {len(skipped)}")
        for s in skipped:
            print(f"  - {s}")
    if missing_log:
        log_path = work / (xlsx_path.stem + "_unmatched.txt")
        log_path.write_text("\n".join(missing_log))
        print(f"Unmatched proposition list written to: {log_path}")


def _work_dir(xlsx_path: Path) -> Path:
    """`Outputs/Data/` when the form is in `Outputs/`, else beside the form.

    A form pinned somewhere else -- a copy under test, a scratch build -- keeps
    its companions with it rather than having them appear in the project tree.
    """
    try:
        from fpp.paths import OUTPUTS, OUTPUTS_DATA
        if xlsx_path.resolve().parent == OUTPUTS.resolve():
            OUTPUTS_DATA.mkdir(parents=True, exist_ok=True)
            return OUTPUTS_DATA
    except Exception:
        pass
    return xlsx_path.parent


if __name__ == "__main__":
    main()
