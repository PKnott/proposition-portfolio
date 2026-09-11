"""Rebuild the fixture workbook from a previous one, without retraining.

`05_Run` retrains four models before it can write a sheet, which is far too slow
a loop for iterating on layout. Every predicted rate the writer needs is already
recorded in the last workbook's metadata block, so this reads them back out and
calls `write_workbook` directly.

    python scripts/rebuild_workbook_from_output.py [Outputs/predictions_<date>.xlsx]

A layout tool, not a pipeline step: it reproduces presentation from stored
predictions and cannot produce new ones.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd

import fpp
from fpp.clean import load_clean_table
from fpp.spec import STAT_BY_KEY, TARGETS


def read_preds(path: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path)
    contents = wb["Contents"]
    rows = []

    for r in range(3, contents.max_row + 1):
        code = contents.cell(r, 1).value
        if not code:
            continue
        ws = wb[code]
        rec = {
            "fixture_id": f"{code}",
            "date": pd.Timestamp(contents.cell(r, 2).value),
            "league": contents.cell(r, 3).value,
            "home_team": contents.cell(r, 4).value,
            "away_team": contents.cell(r, 5).value,
        }
        # Metadata block: rates sit at D2:E5 (home) and G2:H5 (away), one row per
        # target in `TARGETS` order.
        for i, t in enumerate(TARGETS):
            label = ws.cell(2 + i, 4).value or ""
            assert STAT_BY_KEY[t].display in str(label), f"{code}: expected {t} at row {2 + i}, saw {label!r}"
            rec[f"{t}_home"] = float(ws.cell(2 + i, 5).value)
            rec[f"{t}_away"] = float(ws.cell(2 + i, 8).value)
        rows.append(rec)

    preds = pd.DataFrame(rows)
    league_to_key = {v.display: k for k, v in fpp.config.LEAGUES.items()}
    preds["league_key"] = preds["league"].map(league_to_key)
    missing = preds["league_key"].isna()
    assert not missing.any(), f"unmapped leagues: {sorted(preds.loc[missing, 'league'].unique())}"
    return preds


def _current_predictions() -> Path:
    """The predictions workbook in the `Outputs` root.

    The layout guarantees there is exactly one, which is why this can be a
    default at all -- the old hardcoded date silently stopped resolving the day
    it was archived.
    """
    found = sorted(fpp.paths.OUTPUTS.glob("predictions_*.xlsx"))
    if not found:
        sys.exit("No predictions workbook in Outputs/ -- run 05_Run first, "
                 "or pass one from Outputs/Archive/predictions/.")
    return found[-1]


def main() -> None:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else _current_predictions()
    dst = Path(sys.argv[2] if len(sys.argv) > 2 else fpp.paths.OUTPUTS_DATA / "_layout_check.xlsx")
    dst.parent.mkdir(parents=True, exist_ok=True)

    preds = read_preds(src)
    print(f"recovered {len(preds)} fixtures from {src}")

    tm = load_clean_table()
    ctx = fpp.RunContext.load()
    fpp.report.write_workbook(preds, tm, out_path=dst, dispersion=ctx.dispersion)


if __name__ == "__main__":
    main()
