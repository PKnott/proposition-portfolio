"""Per-proposition invariants: price, edge, dominance, and the two stake splits.

Everything above one proposition -- which mix to back, and how the mixes compare
-- lives in `tests/test_portfolio.py`, matching the split between `fpp.staking`
and `fpp.portfolio`.

The worked example below is the acceptance case for the calculator. Its numbers
were derived independently of this code, so they pin behaviour rather than
record it.

One correction to the spec it came from, which described these three as mutually
non-dominating. On the ``(p, e)`` axes it used they are not -- proposition 3
(P=0.70, E=1.085) beats proposition 1 (P=0.60, E=1.080) on both. On the ``(p, o)``
axes the rule now uses they *are* mutually non-dominating, because proposition 3
pays less. That is the difference the axis change makes, and
`test_the_worked_example_is_mutually_non_dominating_on_p_and_o` pins it.
"""

from __future__ import annotations

import itertools

import numpy as np
import pandas as pd
import pytest

from fpp import staking
from fpp.report.markets import LADDER_LINES, count_markets, goal_markets, ladder_lines
from fpp.spec import STAT_BY_KEY, TARGETS

# E = 1.080, 1.100, 1.085
WE_P = np.array([0.60, 0.55, 0.70])
WE_O = np.array([1.80, 2.00, 1.55])
WE_S = 100.0


def _props(p, o, sheet="X-01", labels=None):
    labels = labels or [f"prop {i}" for i in range(len(p))]
    return pd.DataFrame({"sheet_code": sheet, "label": labels, "p": p, "o": o})


# --- The worked example ----------------------------------------------------


def test_edge_is_the_only_derived_column():
    df = staking.add_edge(_props(WE_P, WE_O))
    assert df["e"].to_numpy() == pytest.approx([1.080, 1.100, 1.085])
    # `m = p * e` is deliberately gone -- it is `o * p**2`, so it counted the
    # model's confidence twice and the price once, and it decided which
    # proposition each match contributed. Nothing replaces it at this level.
    assert "m" not in df.columns


def test_stake_split_matches_the_worked_example():
    s = staking.stake_split(WE_P * WE_O, WE_S)
    assert s == pytest.approx([33.5885, 32.9779, 33.4336], abs=5e-4)
    assert s.sum() == pytest.approx(WE_S, abs=1e-12)


def test_summary_matches_the_worked_example():
    s = staking.stake_split(WE_P * WE_O, WE_S)
    out = staking.summarise(WE_P, WE_O, s, WE_S)
    assert out["expected_return"] == pytest.approx(108.827, abs=5e-4)
    assert out["variance"] == pytest.approx(2517.9, abs=0.05)
    assert out["sd"] == pytest.approx(50.18, abs=5e-3)
    # Three propositions give eight outcomes, and nothing lands between 66 and
    # 112 -- so every threshold at or below 110% reads the same 0.673. The
    # plateau is a property of a lumpy distribution rather than a rule, and it is
    # exactly why `portfolio.threshold_probs_batch` had to be a convolution: a
    # saddlepoint approximation smooths steps like this away.
    plateau = [t for t in staking.THRESHOLDS if t <= 1.10]
    assert len(plateau) == len(staking.THRESHOLDS)
    for t in plateau:
        assert out["thresholds"][t] == pytest.approx(0.673, abs=1e-9)
    # The step is above the reported range now that 1.20 is gone; check it is
    # still there rather than asserting a threshold that no longer exists.
    assert staking.summarise(WE_P, WE_O, s, WE_S, thresholds=(1.20,))[
        "thresholds"][1.20] == pytest.approx(0.33, abs=1e-9)


def test_the_worked_example_is_mutually_non_dominating_on_p_and_o():
    """All three survive, which they did not under the old `(p, e)` rule.

    Proposition 3 (P=0.70, O=1.55) has the highest probability and the shortest
    price; proposition 2 (P=0.55, O=2.00) the lowest and the longest. Neither
    beats the other on both, and nor does 1 against either -- so the frontier is
    all three, and which of them to back is a portfolio question.

    On `(p, e)` proposition 3 beat proposition 1 outright, because `e` carries
    `p` inside it and so counted the same advantage twice. That is the whole
    reason the axes changed, and this is where the difference is visible.
    """
    df = staking.add_edge(_props(WE_P, WE_O, labels=["one", "two", "three"]))
    assert staking.undominated(df)["label"].tolist() == ["one", "two", "three"]
    assert staking.undominated(df, ("p", "e"))["label"].tolist() == ["two", "three"]


# --- Structural properties of the split ------------------------------------


def test_split_equalises_expected_contribution():
    """``s_i * e_i`` is constant -- the defining property of the 1/E split."""
    s = staking.stake_split(WE_P * WE_O, WE_S)
    contrib = s * WE_P * WE_O
    assert contrib == pytest.approx(np.full(3, contrib[0]), rel=1e-12)


def test_expected_return_has_the_closed_form_n_S_over_F():
    e = WE_P * WE_O
    F = (1.0 / e).sum()
    s = staking.stake_split(e, WE_S)
    assert staking.expected_return(WE_P, WE_O, s) == pytest.approx(len(e) * WE_S / F, rel=1e-12)


def test_closed_forms_agree_with_full_enumeration():
    s = staking.stake_split(WE_P * WE_O, WE_S)
    _, probs, returns = staking.outcome_paths(WE_P, WE_O, s)
    mean = float((probs * returns).sum())
    assert staking.expected_return(WE_P, WE_O, s) == pytest.approx(mean, rel=1e-12)
    assert staking.variance(WE_P, WE_O, s) == pytest.approx(
        float((probs * (returns - mean) ** 2).sum()), rel=1e-12)


def test_path_probabilities_sum_to_one():
    s = staking.stake_split(WE_P * WE_O, WE_S)
    _, probs, _ = staking.outcome_paths(WE_P, WE_O, s)
    assert probs.sum() == pytest.approx(1.0, rel=1e-12)


@pytest.mark.parametrize("total", [1.0, 100.0, 1e6])
def test_percentages_are_independent_of_the_total_stake(total):
    """Only the cash figures move with S. This is what lets the sheet stay live
    off one Python computation instead of 2**n rows of formulas."""
    ref = staking.summarise(WE_P, WE_O, staking.stake_split(WE_P * WE_O, 100.0), 100.0)
    got = staking.summarise(WE_P, WE_O, staking.stake_split(WE_P * WE_O, total), total)
    assert got["pct_expected_return"] == pytest.approx(ref["pct_expected_return"], rel=1e-12)
    assert got["expected_return"] == pytest.approx(ref["expected_return"] * total / 100.0, rel=1e-12)
    for t in staking.THRESHOLDS:
        assert got["thresholds"][t] == pytest.approx(ref["thresholds"][t], abs=1e-12)


# --- Dominance -------------------------------------------------------------


def test_strictly_worse_on_both_axes_is_dominated():
    """Lower probability *and* a shorter price: nothing can rescue it."""
    df = staking.add_edge(_props([0.60, 0.50], [1.80, 1.70]))
    assert staking.undominated(df)["label"].tolist() == ["prop 0"]


def test_tied_on_p_and_strictly_worse_on_o_is_dominated():
    df = pd.DataFrame({"label": ["A", "B"], "p": [0.60, 0.60], "o": [1.80, 1.70]})
    assert staking.undominated(df)["label"].tolist() == ["A"]


def test_an_exact_tie_on_both_axes_keeps_both():
    df = pd.DataFrame({"label": ["A", "B"], "p": [0.60, 0.60], "o": [1.80, 1.80]})
    assert staking.undominated(df)["label"].tolist() == ["A", "B"]


def test_dominance_uses_full_precision_not_displayed_values():
    """Two propositions identical to 2dp but differing in the 12th are not tied."""
    df = pd.DataFrame({"label": ["A", "B"], "p": [0.60, 0.60 + 1e-12], "o": [1.80, 1.80]})
    assert staking.undominated(df)["label"].tolist() == ["B"]


def test_dominance_keeps_the_pareto_frontier_only():
    rng = np.random.default_rng(0)
    df = pd.DataFrame({"label": [str(i) for i in range(60)],
                       "p": rng.uniform(0.5, 0.95, 60), "o": rng.uniform(1.05, 2.5, 60)})
    kept = staking.undominated(df)
    p, o = kept["p"].to_numpy(), kept["o"].to_numpy()
    for i in range(len(kept)):
        for j in range(len(kept)):
            if i != j:
                assert not ((p[i] >= p[j] and o[i] >= o[j]) and (p[i] > p[j] or o[i] > o[j]))


def test_the_axes_are_a_parameter_not_a_hardcoded_pair():
    """`(p, e)` is still reachable, so the change of default is a decision, not a loss."""
    df = pd.DataFrame({"label": ["A", "B"], "p": [0.60, 0.70], "o": [1.80, 1.55]})
    df = staking.add_edge(df)
    assert staking.undominated(df, ("p", "o"))["label"].tolist() == ["A", "B"]
    assert staking.undominated(df, ("p", "e"))["label"].tolist() == ["B"]


def test_p_o_dominance_prunes_less_than_p_e_did():
    """Not a defect: `p` and `o` are near-inverses, so mutual dominance is rare.

    Pinned because the pruning rate is what sizes the portfolio search, and a
    change here would be felt there rather than caught here.
    """
    rng = np.random.default_rng(11)
    p = rng.uniform(0.15, 0.9, 200)
    df = staking.add_edge(pd.DataFrame(
        {"label": [str(i) for i in range(200)], "p": p,
         "o": (1.0 / p) * rng.uniform(1.0, 1.2, 200)}))
    assert len(staking.undominated(df, ("p", "o"))) > len(staking.undominated(df, ("p", "e")))


# --- Prices ----------------------------------------------------------------
#
# Fixtures below name their book columns *positionally* rather than by key, via
# `_prices`. `BOOKS` is meant to be edited -- the module says adding a book is one
# line -- and hard-coding book keys here turned that one line into a dozen
# failing tests that had nothing to say about the change. Books a test does not
# name are all-NaN, which is what "this book is not pricing that line" already
# means everywhere else.

BOOK1, BOOK2, BOOK3 = staking.BOOK_COLUMNS[:3]
NAME1, NAME2, NAME3 = (staking.BOOKS[b] for b in (BOOK1, BOOK2, BOOK3))


def _prices(*columns):
    """One list per book, in `BOOKS` order; the rest of the books are blank."""
    n = max(len(c) for c in columns)
    out = {b: [np.nan] * n for b in staking.BOOK_COLUMNS}
    out.update(zip(staking.BOOK_COLUMNS, columns))
    return out


def test_best_price_takes_the_best_book_and_tolerates_a_blank():
    df = pd.DataFrame({"sheet_code": "X", "label": ["a", "b", "c"], "p": [0.6, 0.6, 0.6],
                       **_prices([1.80, np.nan, 2.10], [1.95, 2.05, np.nan],
                                 [np.nan, 1.90, 2.05])})
    out = staking.best_price(df)
    # The winner is a different book on each row, and each row has a different
    # book missing -- the max has to skip blanks rather than propagate them.
    assert out["o"].tolist() == [1.95, 2.05, 2.10]


def test_a_proposition_priced_by_no_book_is_dropped():
    df = pd.DataFrame({"sheet_code": "X", "label": ["a", "b"], "p": [0.6, 0.6],
                       **_prices([1.80, np.nan], [1.95, np.nan])})
    assert staking.best_price(df)["label"].tolist() == ["a"]


def test_edge_below_one_is_discarded():
    df = staking.add_edge(_props([0.60, 0.60], [1.80, 1.60]))  # e = 1.08, 0.96
    assert staking.positive_edge(df)["label"].tolist() == ["prop 0"]


# --- Allocation ------------------------------------------------------------


def test_min_variance_closed_form_matches_a_grid_search():
    s = staking.min_variance_alloc(WE_P, WE_O, WE_S)
    best, N = None, 200
    for i in range(N + 1):
        for j in range(N + 1 - i):
            g = np.array([WE_S * i / N, WE_S * j / N, WE_S * (N - i - j) / N])
            v = staking.variance(WE_P, WE_O, g)
            if best is None or v < best:
                best = v
    assert staking.variance(WE_P, WE_O, s) <= best + 1e-6


def test_min_variance_survives_a_near_certain_proposition():
    """With `MIN_MODEL_P` at zero, `p` near 1 is an ordinary row on the form.

    The weight is `1 / (o**2 p (1-p))`, which diverges there. Unclamped, one
    proposition took the whole stake and every other weight underflowed to zero.
    """
    p = np.array([0.999999999, 0.60, 0.55])
    o = np.array([1.02, 1.80, 2.00])
    s = staking.min_variance_alloc(p, o, WE_S)
    assert np.isfinite(s).all()
    assert s.sum() == pytest.approx(WE_S)
    assert s[0] < WE_S, "a near-certainty swallowed the entire allocation"


# --- Ladder alignment ------------------------------------------------------


def test_goals_use_fixed_lines_and_counts_use_dynamic_ones():
    gm = goal_markets(1.6, 1.2)
    h, a = ladder_lines("goals", gm)
    assert h == a == STAT_BY_KEY["goals"].team_lines
    assert len(h) == 4
    for t in [x for x in TARGETS if x != "goals"]:
        mk = count_markets(t, 12.0, 9.0, dispersion=8.0)
        h, a = ladder_lines(t, mk)
        assert len(h) == len(a) == LADDER_LINES
        assert h == tuple(mk["home_lines"]) and a == tuple(mk["away_lines"])


def test_ladder_lines_is_what_the_workbook_writer_uses():
    """Guards the odds form against drifting from the fixture sheets."""
    import inspect

    from fpp.report import excel
    src = inspect.getsource(excel._write_ladder_block)
    assert "ladder_lines(" in src
    assert "dynamic" not in src


# --- Pipeline --------------------------------------------------------------
#
# `staking.split` is gone; the pipeline from a filled form to a set of bets is
# `portfolio.search`, and `tests/test_portfolio.py` covers it end to end. What
# stays here is the part that is still about one proposition at a time.


def test_calculator_reproduces_the_worked_example_for_three_propositions():
    """The spec's stake and distribution figures, fed straight to the calculator."""
    s = staking.stake_split(WE_P * WE_O, WE_S)
    out = staking.summarise(WE_P, WE_O, s, WE_S)
    assert s == pytest.approx([33.5885, 32.9779, 33.4336], abs=5e-4)
    assert out["expected_return"] == pytest.approx(108.827, abs=5e-4)


def test_thresholds_cover_both_sides_of_break_even():
    """Symmetry is the property, not the particular rungs.

    The tuple has been both five long and three long; what has to hold either way
    is that a table whose purpose is comparing upside reports the downside at the
    same resolution. Pinning the literal made trimming it a test failure that
    said nothing about behaviour.
    """
    assert 1.00 in staking.THRESHOLDS
    assert sum(t < 1.0 for t in staking.THRESHOLDS) == sum(t > 1.0 for t in staking.THRESHOLDS)
    assert staking.THRESHOLDS == tuple(sorted(staking.THRESHOLDS))


# --- Entry order on the form ----------------------------------------------


def _ladder_props():
    """Two fixtures, both teams, every target -- deliberately shuffled."""
    rows = []
    for code, home, away in (("A-01", "Home A", "Away A"), ("A-02", "Home B", "Away B")):
        for target in ("corners", "goals", "sot", "shots"):        # not spec order
            for scope, team in (("away", away), ("home", home)):   # away first
                for line in (2.5, 0.5, 1.5):                       # unsorted
                    rows.append({"sheet_code": code, "target": target, "scope": scope,
                                 "team": team, "line": line, "p": 0.9,
                                 "label": f"{target} - {team} - Over {line}"})
    return pd.DataFrame(rows)


def test_form_is_ordered_for_fast_entry():
    """Fixture, then target in spec order, then home before away, then line up."""
    out = staking.qualifying(_ladder_props(), min_p=0.0)
    for code, g in out.groupby("sheet_code", sort=False):
        seen_targets = list(dict.fromkeys(g["target"]))
        assert seen_targets == list(TARGETS), f"{code}: {seen_targets}"
        for target, gt in g.groupby("target", sort=False):
            assert list(dict.fromkeys(gt["scope"])) == ["home", "away"]
            for scope, gs in gt.groupby("scope", sort=False):
                assert gs["line"].is_monotonic_increasing, f"{code}/{target}/{scope}"


def test_form_keeps_each_fixture_contiguous():
    """The new sort must group by target *inside* a fixture, not across them."""
    out = staking.qualifying(_ladder_props(), min_p=0.0)
    codes = list(dict.fromkeys(out["sheet_code"]))
    assert codes == ["A-01", "A-02"]
    assert out["sheet_code"].tolist() == sorted(out["sheet_code"].tolist(), key=codes.index)


def test_form_order_survives_the_probability_filter():
    """Filtering leaves gaps in the ladders; it must not disturb the order."""
    props = _ladder_props()
    props.loc[props["line"] == 1.5, "p"] = 0.2          # punch a hole in every ladder
    out = staking.qualifying(props, min_p=0.5)
    assert (out["line"] != 1.5).all()
    for _, gt in out.groupby(["sheet_code", "target", "scope"], sort=False):
        assert gt["line"].is_monotonic_increasing


# --- "Not offered" ---------------------------------------------------------


def test_zero_means_not_offered_and_the_other_book_is_used():
    df = pd.DataFrame({"sheet_code": "X", "label": ["a", "b"], "p": [0.6, 0.6],
                       **_prices([staking.NOT_OFFERED, 2.10],
                                 [1.95, staking.NOT_OFFERED])})
    # `read_filled` maps 0 -> NaN; this pins the arithmetic downstream of that.
    df = df.replace(staking.NOT_OFFERED, np.nan)
    out = staking.best_price(df)
    assert out["o"].tolist() == [1.95, 2.10]
    assert out["book"].tolist() == [NAME2, NAME1]


def test_best_price_names_every_book_that_ties_at_the_top():
    """A tie names all of the winners, not "Both".

    With two books "Both" was unambiguous. With three it is not: it cannot say
    whether the third book also pays that price, and the sheet exists to be
    walked into a shop with.
    """
    df = pd.DataFrame({"sheet_code": "X", "label": ["a", "b", "c", "d"], "p": [0.6] * 4,
                       **_prices([1.80, 2.10, 1.90, 1.90],
                                 [1.95, np.nan, 1.90, 1.90],
                                 [np.nan, np.nan, np.nan, 1.90])})
    out = staking.best_price(df)
    assert out["book"].tolist() == [
        NAME2,                                  # one winner
        NAME1,                                  # one winner, others blank
        f"{NAME1} / {NAME2}",                   # two tie
        f"{NAME1} / {NAME2} / {NAME3}",         # three tie
    ]


# --- Form round trip -------------------------------------------------------


def _tiny_form_props():
    return pd.DataFrame({
        "sheet_code": ["A-01"] * 3,
        "date": pd.Timestamp("2026-08-19"), "league": "La Liga",
        "home_team": "Home", "away_team": "Away",
        "label": ["Goals - Home - Over 0.5", "Goals - Home - Over 1.5", "Shots - Away - Over 8.5"],
        "p": [0.90, 0.70, 0.60],
    })


def test_zero_round_trips_through_the_workbook_as_not_offered(tmp_path):
    import openpyxl

    from fpp.report.odds_form import FIRST_DATA_ROW, read_filled, write_odds_form

    path = write_odds_form(_tiny_form_props(), out_path=tmp_path / "form.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["A-01"]
    ws.cell(FIRST_DATA_ROW, 3).value = 0        # "not offered", typed
    ws.cell(FIRST_DATA_ROW, 4).value = 1.85     # priced by one of the others
    ws.cell(FIRST_DATA_ROW, 5).value = 0
    for c in range(3, 3 + len(staking.BOOK_COLUMNS)):
        ws.cell(FIRST_DATA_ROW + 1, c).value = 0    # not offered by any book
    wb.save(path)

    got = read_filled(path)
    assert got[BOOK1].dtype.kind == "f"                 # a real float column, not object
    assert pd.isna(got.loc[0, BOOK1])                   # missing, exactly like a blank
    assert got.loc[0, BOOK2] == 1.85
    priced = staking.best_price(got)
    assert priced["label"].tolist() == ["Goals - Home - Over 0.5"]  # the 0/0 row is gone


def test_a_price_between_zero_and_one_still_raises(tmp_path):
    import openpyxl

    from fpp.report.odds_form import FIRST_DATA_ROW, read_filled, write_odds_form

    path = write_odds_form(_tiny_form_props(), out_path=tmp_path / "form.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["A-01"].cell(FIRST_DATA_ROW, 3).value = 0.85     # 1.85 mistyped
    wb.save(path)
    with pytest.raises(ValueError, match="above 1.00"):
        read_filled(path)


# --- Book attribution ------------------------------------------------------


def test_book_survives_to_the_proposition_table():
    """The output is a list of bets to go and place, so the price needs a source."""
    filled = pd.DataFrame({
        "sheet_code": ["A-01", "A-02", "A-03"], "label": ["one", "two", "three"],
        "p": WE_P, **_prices([1.80, np.nan, 1.55], [np.nan, 2.00, 1.55]),
        "home_team": "H", "away_team": "A",
    })
    from fpp import portfolio as pf

    qualified = pf.qualify(filled)
    names = set(staking.BOOKS.values())
    assert len(qualified) == 3
    for label in qualified["book"]:
        # One book, or a " / "-joined set of books that tied at the top.
        assert set(label.split(" / ")) <= names, label
