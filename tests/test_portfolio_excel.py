"""Guards on the portfolio workbook.

The interesting risk here is not that a cell is the wrong colour, it is that a
**formula points at the wrong cell**. openpyxl writes strings; nothing evaluates
them, so a Match predicate testing the "Sort by" box instead of "Undominated
only" produces a workbook that opens cleanly, filters confidently, and is wrong.
That happened once already -- the four non-numeric query cells were hardcoded at
B15-B18 while the input block ended at B13 -- so the tests below parse the
generated formulas and evaluate them rather than checking they exist.
"""

from __future__ import annotations

import re

import numpy as np
import openpyxl
import pandas as pd
import pytest

from fpp import portfolio as pf
from fpp import staking
from fpp.report import odds_form as px_form
from fpp.report import portfolio_excel as px
from fpp.report.portfolio_excel import write_portfolio_workbook


@pytest.fixture(scope="module")
def result():
    rng = np.random.default_rng(4)
    rows = []
    for j in range(9):
        for i in range(int(rng.integers(1, 4))):
            p = float(rng.uniform(0.25, 0.85))
            rows.append({"sheet_code": f"EV-{j:02d}", "label": f"EV-{j:02d} line {i}",
                         "p": p, "b365": (1.0 / p) * float(rng.uniform(1.03, 1.25)),
                         "home_team": f"Home {j}", "away_team": f"Away {j}"})
    return pf.search(pd.DataFrame(rows), min_legs=1)


@pytest.fixture(scope="module")
def book(result, tmp_path_factory):
    path = tmp_path_factory.mktemp("wb") / "portfolios.xlsx"
    write_portfolio_workbook(result, out_path=path, stake=250.0, source="test.xlsx")
    return openpyxl.load_workbook(path)


# --- Structure --------------------------------------------------------------


def test_every_sheet_is_present(book):
    assert book.sheetnames == ["Contents", "Propositions", "Portfolios", "Query",
                               "Legs", "Stake Calculator"]


def test_portfolios_sheet_matches_the_column_spec(book):
    ws = book["Portfolios"]
    headings = [ws.cell(px.HEADER_ROW, c).value for c in range(1, px.MATCH_COL + 1)]
    assert headings == [c[1] for c in px.COLUMNS] + ["Match"]


def test_undominated_portfolios_are_written_first(book, result):
    ws = book["Portfolios"]
    col = px.COL_KEYS.index("undominated") + 1
    flags = [ws.cell(r, col).value
             for r in range(px.FIRST_DATA_ROW, px.FIRST_DATA_ROW + len(result["scored"]))]
    flags = [f for f in flags if f is not None]
    # A row cap can then only ever drop dominated portfolios.
    assert flags == sorted(flags, key=lambda v: not v)


def test_every_written_portfolio_has_a_match_formula(book, result):
    ws = book["Portfolios"]
    n = min(len(result["scored"]), px.MAX_ROWS)
    for r in (px.FIRST_DATA_ROW, px.FIRST_DATA_ROW + n - 1):
        assert str(ws.cell(r, px.MATCH_COL).value).startswith("=AND(")


# --- The query cells the formulas point at ---------------------------------


def test_the_query_labels_sit_at_the_cells_the_formulas_reference(book):
    """The bug this file exists for: constants drifting from the layout."""
    ws = book["Query"]
    for i, (label, _key, cmp, _default) in enumerate(px.QUERY_ROWS):
        assert ws.cell(px.QUERY_FIRST_ROW + i, 1).value == f"{label}  ({cmp})"
    for cell, label in ((px.SPLIT_CELL, "Split"),
                        (px.UNDOM_CELL, "Undominated only"),
                        (px.SORT_CELL, "Sort by"),
                        (px.DESC_CELL, "Highest first")):
        row = int(cell[1:])
        assert ws.cell(row, 1).value == label, f"{cell} is '{ws.cell(row, 1).value}', not '{label}'"


def test_the_filter_block_reads_the_match_column(book):
    ws = book["Query"]
    formula = ws.cell(px.RESULT_ROW + 1, 1).value
    match_col = openpyxl.utils.get_column_letter(px.MATCH_COL)
    assert f"Portfolios!${match_col}$" in formula
    assert "FILTER(" in formula and "SORT(" in formula
    assert px.SORT_CELL in formula and px.DESC_CELL in formula


def test_the_calculator_reads_the_columns_it_labels(book):
    """Each summary row must index the column whose name it prints."""
    ws = book["Stake Calculator"]
    wanted = {"Expected return": "pct_expected_return", "Standard deviation": "pct_sd",
              "Legs": "legs", "Variance": "variance"}
    # Built from `THRESHOLDS` rather than listed, so trimming the tuple does not
    # leave this asserting on columns the workbook no longer has.
    for t in staking.THRESHOLDS:
        pct = int(round(t * 100))
        label = f"{pct}% of stake" + ("   (% Profitability)" if pct == 100 else "")
        wanted[label] = f"p_over_{pct}"
    seen = 0
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label in wanted:
            col = openpyxl.utils.get_column_letter(px.COL_KEYS.index(wanted[label]) + 1)
            assert f"Portfolios!${col}$" in str(ws.cell(r, 2).value), f"{label} reads the wrong column"
            seen += 1
    assert seen == len(wanted), f"only found {seen} of {len(wanted)} summary rows"


# --- The form must be the current one ---------------------------------------


def _stub_form(path, generated, source, fixtures=("E0-01",)):
    """Minimal odds form -- just the Contents provenance and named sheets."""
    import openpyxl

    wb = openpyxl.Workbook()
    cs = wb.active
    cs.title = "Contents"
    cs.cell(2, 1).value = "Predictions source"; cs.cell(2, 2).value = source
    cs.cell(3, 1).value = "Generated";          cs.cell(3, 2).value = generated
    for f in fixtures:
        wb.create_sheet(f)
    wb.save(path)
    return path


def test_the_newest_form_is_chosen_by_date_not_mtime(tmp_path):
    """Filling a form rewrites it, so an old form topped up later has a newer
    mtime than one generated this morning. Sorting on mtime picks the wrong one."""
    import os, time

    old = _stub_form(tmp_path / "odds_input_2026-08-19.xlsx", "2026-08-19", "predictions_2026-08-19.xlsx")
    new = _stub_form(tmp_path / "odds_input_2026-08-21.xlsx", "2026-08-21", "predictions_2026-08-21.xlsx")
    # Make the OLD file the most recently written one.
    time.sleep(0.01)
    os.utime(old, None)
    assert old.stat().st_mtime > new.stat().st_mtime

    assert px_form.latest_form(tmp_path).name == "odds_input_2026-08-21.xlsx"


def test_a_stale_form_raises_rather_than_being_priced(tmp_path):
    """The failure this guards: `06_Split` had a hardcoded date and silently
    consumed a two-day-old form after the model had been retuned."""
    old = _stub_form(tmp_path / "odds_input_2026-08-19.xlsx", "2026-08-19", "predictions_2026-08-19.xlsx")
    _stub_form(tmp_path / "odds_input_2026-08-21.xlsx", "2026-08-21", "predictions_2026-08-21.xlsx")

    with pytest.raises(ValueError, match="not the newest odds form"):
        px_form.check_form_is_current(old)

    meta = px_form.check_form_is_current(old, strict=False)
    assert meta["is_newest"] is False
    assert meta["generated"] == "2026-08-19"
    assert meta["source"] == "predictions_2026-08-19.xlsx"


def test_the_newest_form_passes_and_reports_its_provenance(tmp_path):
    _stub_form(tmp_path / "odds_input_2026-08-19.xlsx", "2026-08-19", "predictions_2026-08-19.xlsx")
    new = _stub_form(tmp_path / "odds_input_2026-08-21.xlsx", "2026-08-21",
                     "predictions_2026-08-21.xlsx", fixtures=("E0-01", "E0-02"))
    meta = px_form.check_form_is_current(new)
    assert meta["is_newest"] and meta["generated"] == "2026-08-21" and meta["fixtures"] == 2


def test_a_lone_form_is_always_current(tmp_path):
    only = _stub_form(tmp_path / "odds_input_2026-08-21.xlsx", "2026-08-21", "p.xlsx")
    assert px_form.check_form_is_current(only)["is_newest"] is True


# --- Excel will refuse a bare future function -------------------------------

# Anything Excel added after 2007 that this writer might reach for. A bare name in
# the XML invalidates the worksheet part: Excel opens with "We found a problem
# with some content" and offers to repair. It is not a warning and not a #NAME?
# -- the file simply will not open cleanly, which is how a first run of this
# writer shipped.
FUTURE_ONLY = [
    "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "RANDARRAY",
    "XLOOKUP", "XMATCH", "IFNA", "IFS", "SWITCH", "TEXTJOIN", "CONCAT",
    "MAXIFS", "MINIFS", "LET", "LAMBDA", "TEXTSPLIT", "TOCOL", "TOROW",
]


def _formulas(path):
    """Every formula string in the workbook, straight out of the XML.

    Read from the raw parts rather than through openpyxl, because openpyxl hands
    back what it was given -- and what Excel objects to is exactly what is on
    disk.
    """
    import re
    import zipfile

    out = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                x = z.read(name).decode("utf8", "ignore")
                for m in re.finditer(r"<f[^>]*>(.*?)</f>", x, re.S):
                    out.append((name, m.group(1)))
    return out


def test_no_formula_uses_a_bare_future_function(book, tmp_path_factory, result):
    import re

    path = tmp_path_factory.mktemp("fx") / "p.xlsx"
    write_portfolio_workbook(result, out_path=path, stake=100.0)

    offenders = []
    for part, f in _formulas(path):
        for fn in FUTURE_ONLY:
            # A bare call: the name followed by "(", not already namespaced.
            for m in re.finditer(rf"(?<![A-Z0-9_.]){fn}\s*\(", f):
                prefix = f[max(0, m.start() - 12):m.start()]
                if "_xlfn" not in prefix:
                    offenders.append(f"{part}: {fn} in {f[:90]}")
    assert not offenders, "bare future function(s) -- Excel will offer to repair:\n" + "\n".join(offenders)


def test_the_dynamic_array_functions_are_namespaced(book):
    """`FILTER` and `SORT` live in the worksheet sub-namespace, not plain `_xlfn.`."""
    from fpp.report.portfolio_excel import future

    assert future("FILTER") == "_xlfn._xlws.FILTER"
    assert future("SORT") == "_xlfn._xlws.SORT"
    assert future("IFERROR") == "IFERROR", "a 2007 function must not be prefixed"


# --- Evaluating the Match formula ------------------------------------------

_CLAUSE = re.compile(
    r"OR\(ISBLANK\(Query!\$B\$(\d+)\),([A-Z]+)(\d+)(>=|<=)Query!\$B\$(\d+)\)"
)


def _evaluate_match(formula: str, row_values: dict[str, object], query: dict[int, object]) -> bool:
    """Evaluate one generated ``=AND(...)`` against real cell values.

    Deliberately parses the string the writer produced rather than re-deriving
    the predicate: re-deriving would agree with itself no matter which cells the
    workbook actually points at, which is the failure being guarded.
    """
    ok = True
    body = formula[len("=AND("):-1]
    for q_blank, col, _row, cmp, q_cmp in _CLAUSE.findall(body):
        assert q_blank == q_cmp, "ISBLANK and the comparison must test the same cell"
        limit = query.get(int(q_blank))
        if limit is None:
            continue
        v = row_values[col]
        ok &= (v >= limit) if cmp == ">=" else (v <= limit)

    m = re.search(r'OR\(Query!\$B\$(\d+)="",([A-Z]+)\d+=Query!\$B\$\1\)', body)
    assert m, "no split clause in the Match formula"
    want_split = query.get(int(m.group(1)))
    if want_split:
        ok &= row_values[m.group(2)] == want_split

    m = re.search(r"OR\(NOT\(Query!\$B\$(\d+)\),([A-Z]+)\d+\)", body)
    assert m, "no undominated clause in the Match formula"
    if query.get(int(m.group(1))):
        ok &= bool(row_values[m.group(2)])
    return bool(ok)


@pytest.mark.parametrize("query_overrides", [
    {},                                                  # the written defaults
    {"pct_expected_return": 1.05, "p_over_100": 0.60},
    {"legs": 3, "p_over_100": None},                     # min legs, no profitability floor
    {"pct_sd": 0.30},
])
def test_match_column_and_a_pandas_filter_agree(book, result, query_overrides):
    """The whole claim of the Query sheet: one predicate, two renderings.

    The Match formula is parsed and evaluated against each row's real values, and
    the answer is compared with the same filter expressed in pandas. A formula
    pointing at the wrong query cell, or at the wrong data column, diverges here.
    """
    ws = book["Portfolios"]
    written = result["scored"].sort_values(
        ["undominated", "p_over_100"], ascending=[False, False]).head(px.MAX_ROWS).reset_index(drop=True)

    # Query cell values, as Excel would see them after the overrides are typed in.
    query: dict[int, object] = {}
    wanted: list[tuple[str, str, object]] = []
    for i, (_label, key, cmp, default) in enumerate(px.QUERY_ROWS):
        v = query_overrides.get(key, default) if key in query_overrides else default
        query[px.QUERY_FIRST_ROW + i] = v
        if v is not None:
            wanted.append((key, cmp, v))
    query[int(px.SPLIT_CELL[1:])] = ""
    query[int(px.UNDOM_CELL[1:])] = True

    expected = pd.Series(True, index=written.index)
    for key, cmp, v in wanted:
        expected &= (written[key] >= v) if cmp == ">=" else (written[key] <= v)
    expected &= written["undominated"]

    got = []
    for i in range(len(written)):
        r = px.FIRST_DATA_ROW + i
        vals = {openpyxl.utils.get_column_letter(j + 1): ws.cell(r, j + 1).value
                for j in range(len(px.COLUMNS))}
        got.append(_evaluate_match(str(ws.cell(r, px.MATCH_COL).value), vals, query))

    assert got == expected.tolist(), (
        f"Match disagrees with the pandas filter on "
        f"{int((np.array(got) != expected.to_numpy()).sum())} of {len(written)} rows"
    )


# --- Legs and the calculator ------------------------------------------------


def test_legs_are_written_for_every_undominated_portfolio(book, result):
    ws = book["Legs"]
    ids = {ws.cell(r, 1).value for r in range(px.FIRST_DATA_ROW, ws.max_row + 1)}
    ids.discard(None)
    undom = set(result["scored"].loc[result["scored"]["undominated"], "id"])
    assert undom <= ids, f"{len(undom - ids)} undominated portfolios cannot be costed"


def test_leg_stakes_sum_to_one_per_portfolio(book):
    ws = book["Legs"]
    rows = [(ws.cell(r, 1).value, ws.cell(r, 9).value)
            for r in range(px.FIRST_DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value]
    df = pd.DataFrame(rows, columns=["id", "stake"])
    assert df.groupby("id")["stake"].sum().round(6).eq(1.0).all()


def test_the_calculator_defaults_to_a_portfolio_that_has_legs(book):
    ws = book["Stake Calculator"]
    default_id = ws.cell(2, 2).value
    legs = book["Legs"]
    ids = {legs.cell(r, 1).value for r in range(px.FIRST_DATA_ROW, legs.max_row + 1)}
    assert default_id in ids
    assert ws.cell(1, 2).value == 250.0


# --- Degenerate input -------------------------------------------------------


def test_nothing_qualifying_writes_a_placeholder_not_a_crash(tmp_path):
    filled = pd.DataFrame({"sheet_code": ["A", "B"], "label": ["a", "b"],
                           "p": [0.4, 0.4], "b365": [1.5, 1.5],
                           "home_team": "H", "away_team": "A"})
    path = write_portfolio_workbook(pf.search(filled), out_path=tmp_path / "empty.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Info"]
