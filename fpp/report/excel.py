"""The combined workbook: every fixture, every league, every market, one file.

Layout per fixture sheet
------------------------
* header + the four predicted rates
* Goals: scoreline matrix, 1X2, BTTS, and a per-team over ladder at fixed lines
* Shots / Shots on Target / Corners: a per-team over ladder each, `LADDER_LINES`
  lines centred on that team's own predicted rate
* a League Averages panel down the right-hand side -- the reference the colours
  are measured against, visible on the same screen as the thing being coloured

Sheet naming
------------
Excel caps sheet names at 31 characters. The old workbook used
``"{date} {home} v {away}"``, which truncated half the Premier League fixtures --
one sheet lost its opponent entirely (``"17 May Wolverhampton Wanderers "``) --
and adding a league prefix for a five-league workbook would make collisions
certain. Sheets are therefore short deterministic codes (``E0-01``), with a
Contents sheet mapping each code to the full fixture and hyperlinking to it.

Colour coding compares each probability to that **league's own** realised rate,
at that **team's own venue**. The model is pooled; the comparison baseline is
not, on either axis.

Match-total over/under is deliberately absent from the fixture blocks. It is a
different bet from the per-team ladders that surround it and reading the two off
one table invited exactly that confusion; the league panel still carries match
rates as reference.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..config import LEAGUES
from ..spec import STAT_BY_KEY, TARGETS, target_spec
from .markets import (
    LeagueBaseline,
    fixture_markets,
    ladder_lines,
    league_baselines,
    matrix_labels,
    safe_odds,
)
from .style import (
    BORDER,
    CENTER,
    FILL_CARD,
    FILL_HEADER,
    FILL_MAX,
    FILL_TABLEH,
    FILL_TITLE,
    FONT_H,
    FONT_TITLE,
    LEFT,
    NUM3,
    ODDS,
    PCT,
    SIGNED_PCT,
    fill_for_delta,
    header,
    kv,
    plain_cell,
    prob_cell,
    table_header,
)

# --- Sheet geometry -------------------------------------------------------
#
# Fixed rather than flowed, so the panel on the right can be written
# independently of how tall the blocks on the left happen to be.

LADDER_COLS = 5  # Line | Over | Odds | Lg Avg | Delta
HOME_COL = 1  # A
AWAY_COL = 7  # G -- one blank column between the two tables
PANEL_COL = 14  # N:O, leaving M blank as a gutter past the 1X2/BTTS block at J:L

MATRIX_ROW = 7
BLOCK_GAP = 2


def sheet_codes(preds: pd.DataFrame) -> list[str]:
    """``E0-01``, ``SP1-07`` ... unique, short, and stable within a run."""
    counts: dict[str, int] = {}
    codes = []
    for lk in preds["league_key"]:
        prefix = LEAGUES[lk].prefix
        counts[prefix] = counts.get(prefix, 0) + 1
        codes.append(f"{prefix}-{counts[prefix]:02d}")
    return codes


# --- Blocks ---------------------------------------------------------------


def _write_matrix(ws, start_row: int, mk: dict, home: str, away: str, base: LeagueBaseline) -> int:
    """Scoreline matrix + 1X2/BTTS. Returns the next free row."""
    labels = matrix_labels()
    M = mk["matrix"]

    header(ws, start_row, 1, 1 + len(labels), "Goals - Scoreline Probability Matrix")
    hdr = start_row + 1
    ws.cell(hdr, 1).fill = FILL_TABLEH
    ws.cell(hdr, 1).border = BORDER
    table_header(ws, hdr, 2, [f"{away} {lab}" for lab in labels])

    mi, mj = mk["most_likely"]
    for i, lab in enumerate(labels):
        r = hdr + 1 + i
        a = ws.cell(r, 1)
        a.value = f"{home} {lab}"
        a.fill = FILL_TABLEH
        a.font = FONT_H
        a.alignment = CENTER
        a.border = BORDER
        for j in range(len(labels)):
            c = plain_cell(ws, r, 2 + j, float(M[i, j]), PCT)
            if (i, j) == (mi, mj):
                c.fill = FILL_MAX
                c.font = FONT_H

    # 1X2 / BTTS to the right of the matrix.
    col = 2 + len(labels) + 1
    header(ws, start_row, col, col + 2, "1X2 / BTTS")
    rows = [
        ("Home", mk["p_home"], base.p_home),
        ("Draw", mk["p_draw"], base.p_draw),
        ("Away", mk["p_away"], base.p_away),
        ("BTTS Yes", mk["p_btts_yes"], base.p_btts_yes),
        ("BTTS No", mk["p_btts_no"], base.p_btts_no),
    ]
    for i, (lab, p, avg) in enumerate(rows):
        r = hdr + 1 + i
        plain_cell(ws, r, col, lab, fill=FILL_CARD, bold=True)
        prob_cell(ws, r, col + 1, p, avg)
        plain_cell(ws, r, col + 2, safe_odds(p), ODDS)

    return hdr + len(labels) + BLOCK_GAP


def _write_one_ladder(ws, start_row: int, col: int, title: str, target: str,
                      lines, probs, base: LeagueBaseline, scope: str) -> None:
    """One team's over ladder: Line | Over | Odds | Lg Avg | Delta.

    ``scope`` routes the baseline to that team's own venue, which is the whole
    point of the block -- an away side priced against home-and-away pooled rates
    reads as underperforming even when it is not.
    """
    header(ws, start_row, col, col + LADDER_COLS - 1, title)
    table_header(ws, start_row + 1, col, ["Line", "Over", "Odds", "Lg Avg", "Δ vs Lg"])

    for i, line in enumerate(lines):
        r = start_row + 2 + i
        p = float(probs[i])
        avg = base.over_rate(target, scope, line)

        plain_cell(ws, r, col, f"Over {line}", fill=FILL_CARD, bold=True)
        prob_cell(ws, r, col + 1, p, avg)
        plain_cell(ws, r, col + 2, safe_odds(p), ODDS)
        plain_cell(ws, r, col + 3, None if avg != avg else avg, PCT)
        plain_cell(ws, r, col + 4, None if avg != avg else p - avg, SIGNED_PCT)


def _write_ladder_block(ws, start_row: int, target: str, mk: dict,
                        home: str, away: str, base: LeagueBaseline) -> int:
    """Both teams' ladders, side by side. Returns the next free row.

    Which lines to price is `markets.ladder_lines`' decision, not this
    function's: the odds capture form has to produce the same set, and one
    definition is what keeps the two workbooks aligned row for row.
    """
    s = target_spec(target)
    home_lines, away_lines = ladder_lines(target, mk)

    for col, team, lines, scope in (
        (HOME_COL, home, home_lines, "home"),
        (AWAY_COL, away, away_lines, "away"),
    ):
        probs = [mk[f"{scope}_over_{ln}"] for ln in lines]
        _write_one_ladder(ws, start_row, col, f"{s.display} - {team}", target,
                          lines, probs, base, scope)

    return start_row + 2 + max(len(home_lines), len(away_lines)) + BLOCK_GAP


def _panel_rows(base: LeagueBaseline) -> list[tuple[str, str | None, float | None]]:
    """The reference panel, as (kind, label, value) rows.

    The goals section reproduces the sixteen rows the previous workbook carried,
    in the same order, so a sheet from either version reads the same way. The
    three count families then repeat that shape -- home, away, match -- at their
    own canonical lines.
    """
    rows: list[tuple[str, str | None, float | None]] = [("section", "Goals", None)]
    g = STAT_BY_KEY["goals"]
    rows.append(("row", "BTTS", base.p_btts_yes))
    for line in g.match_lines:
        rows.append(("row", f"Over {line}", base.over_rate("goals", "match", line)))
    rows += [
        ("row", "Home Win", base.p_home),
        ("row", "Draw", base.p_draw),
        ("row", "Away Win", base.p_away),
    ]
    for scope, word in (("home", "Home"), ("away", "Away")):
        for line in g.team_lines:
            rows.append(("row", f"{word} Over {line}", base.over_rate("goals", scope, line)))

    for t in TARGETS:
        if t == "goals":
            continue
        s = STAT_BY_KEY[t]
        rows.append(("section", s.display, None))
        for scope, word in (("home", "Home"), ("away", "Away")):
            for line in s.team_lines:
                rows.append(("row", f"{word} Over {line}", base.over_rate(t, scope, line)))
        for line in s.match_lines:
            rows.append(("row", f"Match Over {line}", base.over_rate(t, "match", line)))
    return rows


def _write_league_panel(ws, start_row: int, base: LeagueBaseline, league: str) -> int:
    """The uncoloured reference block down the right-hand side.

    Deliberately uncoloured: this is the zero of the delta scale, and shading it
    would invite reading it as another prediction.
    """
    col = PANEL_COL
    span = (f"League Averages - {league}  "
            f"({len(base.seasons)} seasons to {base.seasons[-1]}, {base.n_matches:,} matches)"
            if base.seasons else f"League Averages - {league}")
    header(ws, start_row, col, col + 1, span)

    r = start_row + 1
    for kind, label, value in _panel_rows(base):
        if kind == "section":
            header(ws, r, col, col + 1, label)
        else:
            plain_cell(ws, r, col, label, fill=FILL_CARD, bold=True).alignment = LEFT
            plain_cell(ws, r, col + 1, None if value is None or value != value else value, PCT)
        r += 1
    return r


def _write_fixture_sheet(ws, row: pd.Series, mk: dict, base: LeagueBaseline) -> None:
    home, away = row["home_team"], row["away_team"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    t = ws.cell(1, 1)
    t.value = f"{home} vs {away}  -  {pd.Timestamp(row['date']).date()}  -  {row['league']}"
    t.fill = FILL_TITLE
    t.font = FONT_TITLE
    t.alignment = LEFT

    kv(ws, 2, 1, "League", row["league"])
    kv(ws, 3, 1, "Date", str(pd.Timestamp(row["date"]).date()))
    for i, tgt in enumerate(TARGETS):
        s = STAT_BY_KEY[tgt]
        kv(ws, 2 + i, 4, f"{s.display} ({home})", float(row[f"{tgt}_home"]), NUM3)
        kv(ws, 2 + i, 7, f"{s.display} ({away})", float(row[f"{tgt}_away"]), NUM3)

    r = _write_matrix(ws, MATRIX_ROW, mk["goals"], home, away, base)
    r = _write_ladder_block(ws, r, "goals", mk["goals"], home, away, base)
    for tgt in TARGETS:
        if tgt == "goals":
            continue
        r = _write_ladder_block(ws, r, tgt, mk[tgt], home, away, base)

    _write_league_panel(ws, MATRIX_ROW, base, str(row["league"]))

    for col in range(1, PANEL_COL + 2):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions[get_column_letter(PANEL_COL)].width = 22
    ws.column_dimensions[get_column_letter(PANEL_COL + 1)].width = 12
    ws.freeze_panes = "A2"


def _write_contents(ws, preds: pd.DataFrame, codes: list[str]) -> None:
    ws.merge_cells("A1:F1")
    t = ws.cell(1, 1)
    t.value = f"Predictions - {len(preds)} fixtures across {preds['league_key'].nunique()} leagues"
    t.fill = FILL_TITLE
    t.font = FONT_TITLE
    t.alignment = LEFT

    for j, h in enumerate(["Sheet", "Date", "League", "Home", "Away", "Exp. Goals"]):
        c = ws.cell(2, 1 + j)
        c.value = h
        c.fill = FILL_HEADER
        c.font = FONT_H
        c.alignment = CENTER
        c.border = BORDER

    for i, (code, row) in enumerate(zip(codes, preds.to_dict("records"))):
        r = 3 + i
        link = ws.cell(r, 1)
        link.value = code
        link.hyperlink = f"#'{code}'!A1"
        link.font = Font(color="0563C1", underline="single")
        link.alignment = CENTER
        link.border = BORDER
        for j, v in enumerate([
            str(pd.Timestamp(row["date"]).date()), row["league"],
            row["home_team"], row["away_team"],
            round(float(row["goals_home"]) + float(row["goals_away"]), 2),
        ]):
            plain_cell(ws, r, 2 + j, v)

    for col, w in zip("ABCDEF", (12, 12, 18, 24, 24, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"


_EMPTY_BASELINE = LeagueBaseline(
    league_key="", seasons=(), n_matches=0,
    p_home=float("nan"), p_draw=float("nan"), p_away=float("nan"),
    p_btts_yes=float("nan"), p_btts_no=float("nan"), surv={},
)


def write_workbook(
    preds: pd.DataFrame,
    team_matches: pd.DataFrame,
    out_path: Path | None = None,
    dispersion: dict | None = None,
) -> Path:
    """One workbook covering every fixture in the range, across all five leagues."""
    from ..paths import OUTPUTS, publish

    out_path = Path(out_path) if out_path else OUTPUTS / f"predictions_{dt.date.today().isoformat()}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    if preds.empty:
        with pd.ExcelWriter(out_path, engine="openpyxl") as w:
            pd.DataFrame({"Message": ["No fixtures in the selected date range."]}).to_excel(
                w, sheet_name="Info", index=False)
        publish(out_path, "predictions")
        print(f"No fixtures -- wrote placeholder to {out_path}")
        return out_path

    bases = league_baselines(team_matches)
    codes = sheet_codes(preds)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="Contents", index=False)
        _write_contents(writer.book["Contents"], preds, codes)

        for code, row in zip(codes, [preds.iloc[i] for i in range(len(preds))]):
            pd.DataFrame().to_excel(writer, sheet_name=code, index=False)
            mk = fixture_markets(row, dispersion)
            base = bases.get(row["league_key"], _EMPTY_BASELINE)
            _write_fixture_sheet(writer.book[code], row, mk, base)

    publish(out_path, "predictions")
    print(f"Wrote {len(preds)} fixture sheets -> {out_path}")
    return out_path


__all__ = ["write_workbook", "sheet_codes", "fill_for_delta"]
