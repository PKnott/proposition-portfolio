"""The calibration workbook: metrics, bands and figures in one openable file.

The evaluation used to leave calibration inside the notebook -- plots that
existed only while the kernel was alive, and an ECE in a plot title that nothing
recorded. This writes the whole thing down: a summary across every target, a
sheet per target with the per-band verdicts, and the figures embedded so the
file is readable without rerunning anything.

Colour comes from the same vocabulary as the predictions workbook (`style.py`),
so green means the same direction in both.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from .style import (
    FILL_CARD,
    FILL_G1,
    FILL_N0,
    FILL_R1,
    FILL_R2,
    FILL_TITLE,
    FONT_H,
    FONT_TITLE,
    LEFT,
    PCT,
    SIGNED_PCT,
    header,
    plain_cell,
    table_header,
)

VERDICT_FILL = {
    "TRUST": FILL_G1,
    "CAUTION": FILL_N0,
    "AVOID": FILL_R2,
    "calibrated": FILL_G1,
    "over-predicts": FILL_R1,
    "under-predicts": FILL_R1,
    "insufficient": FILL_CARD,
}

# The headline table. Deliberately short: this is the answer, and the twenty-odd
# diagnostic columns underneath it are the working.
VERDICT_COLS = [
    ("market", "Market", None),
    ("metric", "Scope", None),
    ("n", "N", "#,##0"),
    ("predicted", "Predicted", PCT),
    ("actual", "Actual", PCT),
    ("ece", "ECE", "0.000"),
    ("slope", "Slope", "0.00"),
    ("bands_off", "Bands off", "0"),
    ("verdict", "Verdict", None),
    ("reading", "What is wrong with it", None),
]

ROLLUP_COLS = [
    ("target", "Target", None),
    ("metric", "Scope", None),
    ("trust", "Trust", "0"),
    ("caution", "Caution", "0"),
    ("avoid", "Avoid", "0"),
    ("mean_ece", "Mean ECE", "0.000"),
    ("mean_slope", "Mean slope", "0.00"),
    ("trusted_markets", "Trusted", None),
    ("avoid_markets", "Avoid", None),
]

SUMMARY_COLS = [
    ("segment", "Segment", None),
    ("line", "Line", "0.0"),
    ("n", "N", "#,##0"),
    ("n_bins", "Bands", "0"),
    ("base_rate", "Actual", PCT),
    ("pred_mean", "Predicted", PCT),
    ("ece", "ECE", "0.000"),
    ("mce", "MCE", "0.000"),
    ("brier", "Brier", "0.0000"),
    ("brier_reliability", "  reliability", "0.0000"),
    ("brier_resolution", "  resolution", "0.0000"),
    ("brier_uncertainty", "  uncertainty", "0.0000"),
    ("logloss", "Log loss", "0.0000"),
    ("auc", "AUC", "0.000"),
    ("slope", "Slope", "0.000"),
    ("slope_se", "Slope SE", "0.000"),
    ("intercept", "Intercept", "0.000"),
    ("worst_band_gap", "Worst band", SIGNED_PCT),
    ("n_bands_off", "Bands off", "0"),
    ("confidence", "Reading", None),
    ("verdict", "Verdict", None),
]

BIN_COLS = [
    ("line", "Line", "0.0"),
    ("segment", "Segment", None),
    ("bin", "Band", "0"),
    ("lo", "From", PCT),
    ("hi", "To", PCT),
    ("n", "N", "#,##0"),
    ("pred_mean", "Predicted", PCT),
    ("obs_freq", "Observed", PCT),
    ("gap", "Gap", SIGNED_PCT),
    ("ci_lo", "CI low", PCT),
    ("ci_hi", "CI high", PCT),
    ("verdict", "Verdict", None),
]


def _title(ws, text: str, width: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(1, 1)
    c.value = text
    c.fill = FILL_TITLE
    c.font = FONT_TITLE
    c.alignment = LEFT


def _write_table(ws, start_row: int, frame: pd.DataFrame, cols: list[tuple[str, str, str | None]],
                 caption: str | None = None) -> int:
    """A bordered table with the verdict column colour-coded. Returns next free row."""
    r = start_row
    if caption:
        header(ws, r, 1, len(cols), caption)
        r += 1

    table_header(ws, r, 1, [label for _, label, _ in cols])
    r += 1

    for _, row in frame.iterrows():
        for j, (key, _, fmt) in enumerate(cols):
            v = row.get(key)
            if isinstance(v, float) and not np.isfinite(v):
                v = None
            cell = plain_cell(ws, r, 1 + j, v, fmt)
            if key == "verdict" and v in VERDICT_FILL:
                cell.fill = VERDICT_FILL[v]
                cell.font = FONT_H
        r += 1
    return r + 1


def _write_summary_sheet(ws, verdicts: pd.DataFrame, rollup: pd.DataFrame) -> None:
    """Which markets can be trusted -- the sheet to read first."""
    cols = [("target", "Target", None)] + VERDICT_COLS
    _title(ws, "Calibration - which markets can be trusted (held-out test seasons)", len(cols))

    r = _write_table(ws, 3, rollup, ROLLUP_COLS, "Verdict count by target and scope")
    _write_table(ws, r, verdicts.sort_values(["target", "metric", "line"]), cols,
                 "Every market, with what is wrong with it")

    widths = {"Market": 13, "Scope": 9, "Verdict": 11, "What is wrong with it": 78,
              "Trusted": 34, "Avoid": 34, "Target": 11}
    for j, (_, label, _) in enumerate(cols):
        ws.column_dimensions[ws.cell(2, 1 + j).column_letter].width = widths.get(label, 12)
    ws.freeze_panes = "A4"


def _write_target_sheet(ws, target: str, summary: pd.DataFrame, bins: pd.DataFrame,
                        figures: list[Path], verdicts: pd.DataFrame) -> None:
    _title(ws, f"{target} - calibration detail", len(SUMMARY_COLS))

    v = verdicts[verdicts["target"] == target].sort_values(["metric", "line"])
    r = _write_table(ws, 3, v, VERDICT_COLS, f"{target} - verdict per market") if not v.empty else 3

    for metric in ("team", "total"):
        s = summary[(summary["target"] == target) & (summary["metric"] == metric)]
        if s.empty:
            continue
        r = _write_table(ws, r, s[s["segment"] == "all"].sort_values("line"), SUMMARY_COLS,
                         f"{metric} - pooled")
        seg = s[s["segment"] != "all"].sort_values(["segment", "line"])
        if not seg.empty:
            r = _write_table(ws, r, seg, SUMMARY_COLS, f"{metric} - by league and venue")

        b = bins[(bins["target"] == target) & (bins["metric"] == metric) & (bins["segment"] == "all")]
        if not b.empty:
            r = _write_table(ws, r, b.sort_values(["line", "bin"]), BIN_COLS,
                             f"{metric} - per probability band (pooled)")

    _embed(ws, figures, r)

    for j in range(len(SUMMARY_COLS)):
        ws.column_dimensions[ws.cell(2, 1 + j).column_letter].width = 13
    ws.column_dimensions["J"].width = 70  # the reading column
    ws.freeze_panes = "A3"


def _embed(ws, figures: list[Path], row: int) -> None:
    """Drop the PNGs in below the tables. Silently skipped if Pillow is absent."""
    try:
        from openpyxl.drawing.image import Image
    except ImportError:  # pragma: no cover - openpyxl always ships the module
        return
    for p in figures:
        try:
            img = Image(str(p))
        except Exception:
            # Pillow missing or the file unreadable -- the tables are the payload,
            # so a missing picture must not cost the whole workbook.
            continue
        img.anchor = f"A{row}"
        ws.add_image(img)
        row += max(24, int(img.height / 19) + 2)


def write_calibration_workbook(
    summary: pd.DataFrame,
    bins: pd.DataFrame,
    figures: dict[str, list[Path]] | None = None,
    out_path: Path | None = None,
) -> Path:
    """Write ``Outputs/Evaluation/calibration.xlsx``."""
    from ..calibration import line_verdicts, summarise_by_target
    from ..paths import OUTPUTS

    out_path = Path(out_path) if out_path else OUTPUTS / "Evaluation" / "calibration.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    figures = figures or {}

    verdicts = line_verdicts(summary, bins)
    rollup = summarise_by_target(verdicts)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="Summary", index=False)
        _write_summary_sheet(writer.book["Summary"], verdicts, rollup)

        for target in summary["target"].drop_duplicates():
            pd.DataFrame().to_excel(writer, sheet_name=target, index=False)
            _write_target_sheet(writer.book[target], target, summary, bins,
                                figures.get(target, []), verdicts)

    print(f"Wrote calibration workbook -> {out_path}")
    return out_path


__all__ = ["write_calibration_workbook"]
