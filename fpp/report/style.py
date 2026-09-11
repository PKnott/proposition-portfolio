"""The shared visual vocabulary for every workbook this project writes.

Two workbooks now carry the same delta scale -- the fixture predictions and the
calibration report -- and a colour that means "well above baseline" on one sheet
has to mean the same thing on the other. Keeping the fills in one module is what
makes that true by construction rather than by two files happening to agree.
"""

from __future__ import annotations

import numpy as np
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

THIN = Side(style="thin", color="3A3A3A")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FILL_TITLE = PatternFill("solid", fgColor="1F4E79")
FILL_HEADER = PatternFill("solid", fgColor="D9ECFF")
FILL_CARD = PatternFill("solid", fgColor="F2F6FF")
FILL_TABLEH = PatternFill("solid", fgColor="EAF5EA")
FILL_MAX = PatternFill("solid", fgColor="FFF59D")

# Delta scale vs the league's own rate.
FILL_G3 = PatternFill("solid", fgColor="1B5E20")
FILL_G2 = PatternFill("solid", fgColor="2E7D32")
FILL_G1 = PatternFill("solid", fgColor="C8E6C9")
FILL_N0 = PatternFill("solid", fgColor="FFFFFF")
FILL_R1 = PatternFill("solid", fgColor="FFCDD2")
FILL_R2 = PatternFill("solid", fgColor="E57373")
FILL_R3 = PatternFill("solid", fgColor="B71C1C")

FONT_TITLE = Font(color="FFFFFF", bold=True, size=14)
FONT_H = Font(bold=True)
# The dark ends of the scale need light text; the old workbook left them black
# on dark green/red, which was unreadable exactly when it mattered most.
FONT_ON_DARK = Font(color="FFFFFF", bold=True)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

PCT = "0.00%"
SIGNED_PCT = "+0.00%;-0.00%"
ODDS = "0.00"
NUM3 = "0.000"


def fill_for_delta(delta: float | None) -> tuple[PatternFill, Font | None]:
    """Map a probability gap to a fill, and to a font when the fill is dark."""
    if delta is None or (isinstance(delta, float) and np.isnan(delta)):
        return FILL_N0, None
    if delta >= 0.15:
        return FILL_G3, FONT_ON_DARK
    if delta >= 0.08:
        return FILL_G2, FONT_ON_DARK
    if delta >= 0.03:
        return FILL_G1, None
    if delta <= -0.15:
        return FILL_R3, FONT_ON_DARK
    if delta <= -0.08:
        return FILL_R2, FONT_ON_DARK
    if delta <= -0.03:
        return FILL_R1, None
    return FILL_N0, None


# --- Cell helpers ---------------------------------------------------------


def header(ws, row: int, col_from: int, col_to: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=col_from, end_row=row, end_column=col_to)
    c = ws.cell(row, col_from)
    c.value = text
    c.fill = FILL_HEADER
    c.font = FONT_H
    c.alignment = CENTER
    for cc in range(col_from, col_to + 1):
        ws.cell(row, cc).border = BORDER


def kv(ws, row: int, col: int, label: str, value, fmt: str | None = None) -> None:
    a = ws.cell(row, col)
    a.value = label
    a.font = FONT_H
    a.alignment = LEFT
    a.border = BORDER
    a.fill = FILL_CARD
    b = ws.cell(row, col + 1)
    b.value = value
    b.alignment = CENTER
    b.border = BORDER
    if fmt:
        b.number_format = fmt


def table_header(ws, row: int, col: int, labels: list[str]) -> None:
    """One row of column headings in the table-header fill."""
    for j, text in enumerate(labels):
        c = ws.cell(row, col + j)
        c.value = text
        c.fill = FILL_TABLEH
        c.font = FONT_H
        c.alignment = CENTER
        c.border = BORDER


def plain_cell(ws, row: int, col: int, value, fmt: str | None = None,
               fill: PatternFill | None = None, bold: bool = False):
    """A bordered, centred cell with no conditional colouring."""
    c = ws.cell(row, col)
    c.value = value
    c.alignment = CENTER
    c.border = BORDER
    if fmt:
        c.number_format = fmt
    if fill is not None:
        c.fill = fill
    if bold:
        c.font = FONT_H
    return c


def prob_cell(ws, row: int, col: int, p: float, avg: float | None = None):
    """A probability, coloured by its gap to ``avg`` when one is supplied."""
    c = ws.cell(row, col)
    c.value = float(p)
    c.number_format = PCT
    c.border = BORDER
    c.alignment = CENTER
    delta = None if avg is None or np.isnan(avg) else float(p) - float(avg)
    fill, font = fill_for_delta(delta)
    c.fill = fill
    if font:
        c.font = font
    return c


__all__ = [
    "THIN", "BORDER",
    "FILL_TITLE", "FILL_HEADER", "FILL_CARD", "FILL_TABLEH", "FILL_MAX",
    "FILL_G3", "FILL_G2", "FILL_G1", "FILL_N0", "FILL_R1", "FILL_R2", "FILL_R3",
    "FONT_TITLE", "FONT_H", "FONT_ON_DARK",
    "CENTER", "LEFT",
    "PCT", "SIGNED_PCT", "ODDS", "NUM3",
    "fill_for_delta", "header", "kv", "table_header", "plain_cell", "prob_cell",
]
